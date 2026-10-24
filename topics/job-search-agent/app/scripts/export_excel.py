"""Export the researched feed, source evidence and blank application tracking fields."""
import csv
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT=Path(__file__).resolve().parents[1]
data=json.loads((ROOT/'public/data/jobs.json').read_text())
sources=json.loads((ROOT/'public/data/sources.json').read_text())
wb=Workbook();ws=wb.active;ws.title='Internships'
columns={'company':'Company','title':'Role','roleType':'Role type','location':'Location','workMode':'Work mode','payMin':'Minimum INR/month','payMax':'Maximum INR/month','payVerified':'Pay confirmed in original','status':'Status','nextAction':'Next action','dueDate':'Due date','interviewAt':'Interview time','appliedAt':'Applied date','sourceType':'Platform','sourceUrl':'Original source','applicationUrl':'Application link','email':'Application email','applicationMethod':'Application route','evidence':'Evidence / caveats','requirements':'Requirements','applicationInstructions':'Application instructions','postedLabel':'Posted','firstSeen':'First discovered','checkedAt':'Last checked','notes':'Personal notes'}
ws.append(list(columns.values()))
for job in data['jobs']:
 row=[]
 for key in columns:
  value=job.get(key,'')
  if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@')):value="'"+value
  row.append(value)
 ws.append(row)
 for column in (15,16):
  cell=ws.cell(ws.max_row,column)
  if isinstance(cell.value,str) and cell.value.startswith(('https://','mailto:')):cell.hyperlink=cell.value;cell.style='Hyperlink'
status=DataValidation(type='list',formula1='"discovered,shortlisted,on_hold,applied,assessment,interview,offer,closed"');ws.add_data_validation(status);status.add('I2:I10000')
ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
for col in ws.columns:
 letter=col[0].column_letter
 ws.column_dimensions[letter].width=55 if letter in ['M','N','Q','R','S'] else 27
for row in ws.iter_rows(min_row=2):
 for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
 ws.row_dimensions[row[0].row].height=95
ss=wb.create_sheet('Sources');ss.append(['Company','URL','Type','Fetch result','Last fetch','Notes'])
for s in sources['sources']:ss.append([s['company'],s['url'],s['kind'],s.get('lastStatus',''),s.get('lastChecked',''),s.get('message','')])
ss.freeze_panes='A2';ss.auto_filter.ref=ss.dimensions
company_path=ROOT/'public/data/companies.json'
company_records=[]
if company_path.exists():
 cs=wb.create_sheet('Company watchlist')
 cs.append(['Company','Suggested category','Why this target','Pay evidence type','Reported role','Range minimum','Range maximum','Original period','Confidence','Limitations','Careers URL','Pay source','Observed date'])
 company_records=json.loads(company_path.read_text()).get('companies',[])
 for company in company_records:
  benchmark=company.get('payBenchmark') or {}
  cs.append([company['name'],company['tier'],company['rationale'],benchmark.get('kind','Not established'),benchmark.get('role',''),benchmark.get('min'),benchmark.get('max'),benchmark.get('period',''),benchmark.get('confidence','Unconfirmed'),company.get('payNote',''),company.get('careersUrl',''),benchmark.get('sourceUrl',''),company.get('researchedAt','')])
  for column in (11,12):
   cell=cs.cell(cs.max_row,column)
   if isinstance(cell.value,str) and cell.value.startswith('https://'):cell.hyperlink=cell.value;cell.style='Hyperlink'
 cs.freeze_panes='C2';cs.auto_filter.ref=cs.dimensions
 tiers=DataValidation(type='list',formula1='"Dream,Workable,Possible"');cs.add_data_validation(tiers);tiers.add('B2:B10000')
directory=wb.create_sheet('Company directory')
directory.append(['Company','Category','Company type','Sector','Locations / remote limits','Opening status','Likely roles (inferred)','Careers URL','Public email','Contact page','Contact kind','Notes','Checked'])
requirements=wb.create_sheet('Role requirements')
requirements.append(['Company','Role','Seniority','Location','Role source','Evidence / access caveat','Requirement','Self assessment','Portfolio / resume evidence'])
for company in company_records:
 contact=company.get('contact') or {}
 directory.append([company['name'],company['tier'],company.get('companyType','Unclassified'),company.get('focus',''),company.get('locationNote',''),company.get('openingStatus','No current opening verified'),'; '.join(company.get('likelyRoles',[])),company.get('careersUrl',''),contact.get('email',''),contact.get('url',''),contact.get('kind',''),company.get('notes',''),company.get('researchedAt','')])
 for role in company.get('roles',[]):
  reqs=role.get('requirements') or ['No detailed requirements captured; open source']
  for requirement in reqs:
   text=requirement if isinstance(requirement,str) else requirement.get('text','')
   requirements.append([company['name'],role['title'],role.get('seniority',''),role.get('location',''),role.get('url',''),role.get('evidenceStatus',''),text,'Unassessed',''])
for sheet in (directory,requirements):
 sheet.freeze_panes='C2';sheet.auto_filter.ref=sheet.dimensions
 for row in sheet.iter_rows(min_row=2):
  for cell in row:
   if isinstance(cell.value,str):
    if cell.value.lstrip().startswith(('=','+','-','@')):cell.value="'"+cell.value
    elif cell.value.startswith('https://'):cell.hyperlink=cell.value;cell.style='Hyperlink'
assessment=DataValidation(type='list',formula1='"Unassessed,Met,Partly met,Gap"');requirements.add_data_validation(assessment);assessment.add('H2:H10000')
notes=wb.create_sheet('Read me')
for row in [
 ['Design internship search','Research snapshot: '+data['meta']['researchedAt']],
 ['Priority','Remote India, then Gujarat/Maharashtra/Goa/southern India.'],
 ['Stipend','Gujarat minimum INR20,000/month; elsewhere INR30,000+, INR25,000 conditional. Remote floor provisionally INR30,000.'],
 ['Evidence','Only employer-stated minimums count as confirmed. Undisclosed-pay roles stay on hold. Glassdoor and other company reports are historical context, never a current offer. Annual figures stay annual.'],
 ['Company groups','Dream / Workable / Possible are editable suggestions, not hiring probabilities. The company watchlist does not assert current vacancies or offers.'],
 ['Personal fit','No resume or portfolio assessed. Requirement rows begin Unassessed. Met/partial/gap is an evidence checklist, not hiring odds.'],
 ['Blacklist','Browser-local blacklist hides recommendations, including future matching feed items. Saved application history is retained. Use JSON backup to transfer preferences; daily task cannot read local blacklist.'],
 ['Freshness','Discovery date is not posting date. An accessible listing does not guarantee current hiring.'],
 ['Applications','No applications have been submitted. Status and application fields are for you to update.'],
 ['Board vs Excel','This workbook contains the research snapshot. Browser board changes are exported separately using CSV or JSON backup.'],
 ['Privacy','Initial workbook contains no résumé or personal contact data. Employer application contacts are public listing details.'],
 ['Excluded roles','See the exclusion list below.'],
 *[[x['company'],x['reason']] for x in data['meta']['excluded']]
]:notes.append(row)
for sheet in wb:
 for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='343D2B');cell.alignment=Alignment(wrap_text=True)
 sheet.row_dimensions[1].height=29
 if sheet!=ws:
  for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=35 if col[0].column==1 else 75
  for row in sheet.iter_rows(min_row=2):
   for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
   sheet.row_dimensions[row[0].row].height=42
wb.save(ROOT/'public/data/internships.xlsx')
print('Exported internships.xlsx')

# Public research CSV; private blacklist/preferences/evidence are exported only by the browser.
with (ROOT/'public/data/company-directory.csv').open('w', newline='', encoding='utf-8-sig') as output:
 writer=csv.writer(output)
 writer.writerow(['Company','Category','Type','Sector','Location notes','Opening status','Likely roles (inferred)','Careers URL','Public email','Contact route','Contact kind','Contact caveat','Role','Seniority','Role location','Role URL','Role evidence','Requirements','Personal fit','Pay note','Notes','Sources'])
 for company in company_records:
  contact=company.get('contact') or {}
  for role in company.get('roles') or [{}]:
   row=[company['name'],company['tier'],company.get('companyType','Unclassified'),company.get('focus',''),company.get('locationNote',''),company.get('openingStatus','No current role verified'),'; '.join(company.get('likelyRoles',[])),company.get('careersUrl',''),contact.get('email',''),contact.get('url',''),contact.get('kind',''),contact.get('note',''),role.get('title',''),role.get('seniority',''),role.get('location',''),role.get('url',''),role.get('evidenceStatus',''),'; '.join(r if isinstance(r,str) else r['text'] for r in role.get('requirements',[])),'Unassessed',company.get('payNote','Current pay unconfirmed'),company.get('notes',''),'; '.join(s['url'] for s in company.get('sources',[]))]
   writer.writerow(["'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for v in row])
print('Exported company-directory.csv')
